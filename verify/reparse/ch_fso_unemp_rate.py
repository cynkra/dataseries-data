"""Independent re-parse of ch_fso_unemp_rate (FSO DAM Excel je-d-03.03.01.03), compared cell by
cell with the live API. Does not use the repo's R parser.

Sheet "Monatswerte (1991-...)": row of German month headers like "Jan.91", "Febr.91",
"April.262" (2-digit year, optional trailing footnote digit). Data rows are identified by their
German label in column A:
  Total -> origin=tot, sex=tot; Schweizer/innen -> ch, tot; Ausländer/innen -> ex, tot;
  Männer -> tot, men; Frauen -> tot, wom.
"""
import re
import sys
from datetime import date
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import lib

import openpyxl
import pandas as pd

DS = "ch_fso_unemp_rate"
ORDER = "je-d-03.03.01.03"

MONTHS = {"jan": 1, "feb": 2, "mär": 3, "mar": 3, "apr": 4, "mai": 5, "jun": 6,
          "jul": 7, "aug": 8, "sep": 9, "okt": 10, "nov": 11, "dez": 12}

ROWMAP = {
    "total": ("tot", "tot"),
    "schweizer/innen": ("ch", "tot"),
    "ausländer/innen": ("ex", "tot"),
    "männer": ("tot", "men"),
    "frauen": ("tot", "wom"),
}


def parse_header(h):
    m = re.fullmatch(r"([A-Za-zäÄ]+)\.(\d{2})(\d*)", str(h).strip())
    if not m:
        raise ValueError(f"unparseable header {h!r}")
    yy = int(m.group(2))
    return date(1900 + yy if yy >= 50 else 2000 + yy, MONTHS[m.group(1).lower()[:3]], 1)


def parse(path):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    names = [n for n in wb.sheetnames if n.startswith("Monatswerte (1991")]
    if len(names) != 1:
        raise ValueError(f"expected one 'Monatswerte (1991' sheet, found {wb.sheetnames}")
    rows = list(wb[names[0]].iter_rows(values_only=True))
    hi = next(i for i, r in enumerate(rows)
              if r[1] is not None and re.match(r"[A-Za-zä]+\.\d", str(r[1])))
    dates = [parse_header(h) for h in rows[hi][1:] if h is not None]
    for a, b in zip(dates, dates[1:]):
        if b != date(a.year + (a.month == 12), a.month % 12 + 1, 1):
            raise ValueError(f"month headers not gapless: {a} then {b}")
    out, seen = [], set()
    for r in rows[hi + 1:]:
        lab = str(r[0] or "").strip().lower()
        if lab.startswith("1 ") or lab.startswith("quelle"):
            break
        if lab not in ROWMAP:
            continue
        if lab in seen:
            raise ValueError(f"row label {lab!r} appears twice")
        seen.add(lab)
        o, s = ROWMAP[lab]
        for d, v in zip(dates, r[1:1 + len(dates)]):
            if v is None or (isinstance(v, str) and v.strip() in ("", "X", "...", "…")):
                continue
            out.append({"origin": o, "sex": s, "date": d.isoformat(), "value": float(v)})
    if seen != set(ROWMAP):
        raise ValueError(f"missing rows: {set(ROWMAP) - seen}")
    return pd.DataFrame(out)


def check():
    master, _ = lib.fso_dam_master(ORDER)
    src = parse(lib.download(master, f"{DS}.xlsx"))
    return lib.compare(src, lib.api_csv(DS), ["origin", "sex", "date"])


if __name__ == "__main__":
    only = lib.only_arg(sys.argv[1:])
    if only is None or DS in only:
        v = lib.Verdicts(f"reparse-{DS}")
        v.run(DS, check)
        v.write()
        print(v.items[0]["status"], v.items[0]["summary"])
