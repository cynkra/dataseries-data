"""Independent re-parse of ch_fso_pop (FSO DAM xlsx su-d-01.02.04.05), compared cell by cell
with the live API. Does not use the repo's R parser."""
import re
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import lib

import openpyxl
import pandas as pd

DS = "ch_fso_pop"
ORDER = "su-d-01.02.04.05"

COLS = [
    ("B", "pop_stock_jan",  ["stand am", "1. Januar"]),
    ("C", "live_births",    ["Lebend", "geburten"]),
    ("D", "deaths",         ["Todes", "fälle"]),
    ("E", "birth_surplus",  ["Geburten", "über", "schuss"]),
    ("F", "immigration",    ["Einwanderungen"]),
    ("G", "emigration",     ["Auswanderungen"]),
    ("H", "migration_bal",  ["Wanderungs", "saldo"]),
    ("I", "naturalisation", ["Erwerb des", "Bürgerrechts"]),
    ("J", "adjustments",    ["Bestandes", "bereini"]),
    ("K", "pop_stock_dec",  ["stand am", "31. Dezember"]),
]


def parse(xl):
    ws = openpyxl.load_workbook(xl, data_only=True).worksheets[0]
    for col, code, needles in COLS:
        hdr = " ".join(str(ws[f"{col}{r}"].value or "") for r in range(2, 6))
        for n in needles:
            if n not in hdr:
                raise ValueError(f"header check failed for {col}/{code}: {n!r} not in {hdr!r}")
    rows = []
    for r in range(6, ws.max_row + 1):
        a = ws[f"A{r}"].value
        if a is None:
            continue
        if isinstance(a, float) and a.is_integer():
            a = int(a)
        m = re.fullmatch(r"\s*(\d{4})(?:\s+\d+)?\s*", str(a))
        if not m:
            continue
        for col, code, _ in COLS:
            v = ws[f"{col}{r}"].value
            if v is None or (isinstance(v, str) and v.strip() in ("...", "…", "-", "", "x", "*")):
                continue
            if isinstance(v, str):
                raise ValueError(f"unexpected text {v!r} at {col}{r}")
            rows.append({"item": code, "date": f"{m.group(1)}-01-01", "value": float(v)})
    return pd.DataFrame(rows)


def check():
    master, _ = lib.fso_dam_master(ORDER)
    src = parse(lib.download(master, f"{DS}.xlsx"))
    return lib.compare(src, lib.api_csv(DS), ["item", "date"])


if __name__ == "__main__":
    only = lib.only_arg(sys.argv[1:])
    if only is None or DS in only:
        v = lib.Verdicts(f"reparse-{DS}")
        v.run(DS, check)
        v.write()
        print(v.items[0]["status"], v.items[0]["summary"])
